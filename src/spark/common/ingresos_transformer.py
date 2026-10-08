"""
Modulo de transformacion y limpieza para el dataset de Ingresos Tributarios (SUNAT Cuadro A13).

Transforma el archivo Excel de la SUNAT (cdro_A13.xlsx) que contiene matrices
horizontales con encabezados multinivel y celdas combinadas en dataframes tabulares
con el formato estandarizado:
  - Departamento (str)
  - Periodo (str, ej. '2005_Ene.', '2026_Mar.')
  - Anio (int64)
  - Mes (str, ej. 'Ene.', 'Feb.', 'Dic.')
  - Monto_Recaudado (float64)
"""

import openpyxl
import pandas as pd

MONTH_MAP = {
    "ene": "Ene.",
    "ene.": "Ene.",
    "feb": "Feb.",
    "feb.": "Feb.",
    "mar": "Mar.",
    "mar.": "Mar.",
    "abr": "Abr.",
    "abr.": "Abr.",
    "may": "May.",
    "may.": "May.",
    "jun": "Jun.",
    "jun.": "Jun.",
    "jul": "Jul.",
    "jul.": "Jul.",
    "ago": "Ago.",
    "ago.": "Ago.",
    "sep": "Sep.",
    "sep.": "Sep.",
    "set": "Sep.",
    "set.": "Sep.",
    "oct": "Oct.",
    "oct.": "Oct.",
    "nov": "Nov.",
    "nov.": "Nov.",
    "dic": "Dic.",
    "dic.": "Dic.",
}


def get_dept_name(c1: object | None, c2: object | None) -> str | None:
    """Extrae y normaliza la denominacion de departamento/jurisdiccion."""
    raw = c1 if c1 is not None and str(c1).strip() != "" else c2
    if raw is None:
        return None
    raw_str = str(raw).strip()
    if "provincia constitucional del callao" in raw_str.lower():
        return "Callao"
    return raw_str


def transform_sheet(sheet: openpyxl.worksheet.worksheet.Worksheet) -> pd.DataFrame:
    """Transforma una sola hoja ('cdrA13' o 'cdrA13(Var)') en DataFrame tabular."""
    # 1. Detectar columnas mensuales (año y mes)
    col_meta = {}
    curr_year = None
    for c in range(3, sheet.max_column + 1):
        r4_val = sheet.cell(4, c).value
        r5_val = sheet.cell(5, c).value

        if r4_val is not None:
            try:
                val_int = int(str(r4_val).strip())
                if 2000 <= val_int <= 2030:
                    if curr_year is not None and val_int <= curr_year:
                        # Correccion de cabecera SUNAT cuando repite 2025 para el bloque 2026
                        curr_year = curr_year + 1
                    else:
                        curr_year = val_int
            except ValueError:
                pass

        if r5_val is not None:
            raw_m = str(r5_val).strip().lower()
            if raw_m in MONTH_MAP:
                std_m = MONTH_MAP[raw_m]
                col_meta[c] = (curr_year, std_m)

    # 2. Recorrer filas de departamentos (filas 9 a 37)
    records = []
    for r in range(9, 38):
        c1 = sheet.cell(r, 1).value
        c2 = sheet.cell(r, 2).value
        dept = get_dept_name(c1, c2)
        if not dept:
            continue

        for c, (anio, mes) in col_meta.items():
            val = sheet.cell(r, c).value
            periodo = f"{anio}_{mes}"
            monto = None
            if val is not None:
                try:
                    monto = float(val)
                except (ValueError, TypeError):
                    monto = None

            records.append(
                {
                    "Departamento": dept,
                    "Periodo": periodo,
                    "Anio": int(anio),
                    "Mes": mes,
                    "Monto_Recaudado": monto,
                }
            )

    df = pd.DataFrame(records)
    # Tipado estricto
    df["Departamento"] = df["Departamento"].astype("string")
    df["Periodo"] = df["Periodo"].astype("string")
    df["Anio"] = df["Anio"].astype("int64")
    df["Mes"] = df["Mes"].astype("string")
    df["Monto_Recaudado"] = df["Monto_Recaudado"].astype("float64")
    return df


def process_both_ingresos(
    excel_path_or_file: str | openpyxl.Workbook,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Lee el archivo Excel de Ingresos Tributarios y devuelve (df_tabular, df_var_tabular)."""
    if isinstance(excel_path_or_file, str):
        wb = openpyxl.load_workbook(excel_path_or_file, data_only=True)
    else:
        wb = excel_path_or_file

    sheet_monto = wb["cdrA13"]
    sheet_var = wb["cdrA13(Var)"]

    df_monto = transform_sheet(sheet_monto)
    df_var = transform_sheet(sheet_var)

    return df_monto, df_var
