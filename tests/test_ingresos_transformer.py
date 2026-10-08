import os
import sys
import unittest
from pathlib import Path
import openpyxl

_ROOT = Path(__file__).resolve().parent.parent
_COMMON_PATH = str(_ROOT / "src" / "spark" / "common")
if _COMMON_PATH not in sys.path:
    sys.path.insert(0, _COMMON_PATH)

from ingresos_transformer import process_both_ingresos, transform_sheet

TRASH_EXCEL = Path("/home/gosth/.local/share/Trash/files/cdro_A13.xlsx")


class TestIngresosTransformer(unittest.TestCase):
    def test_mock_workbook_transformation(self):
        """Prueba la transformacion de una hoja sintetica con estructura similar a SUNAT."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "cdrA13"

        # Fila 4: Años
        ws.cell(row=4, column=3, value=2025)
        ws.cell(row=4, column=16, value=2025)  # Typo simulado de SUNAT

        # Fila 5: Meses
        months = ["Ene.", "Feb.", "Mar.", "Abr.", "May.", "Jun.", "Jul.", "Ago.", "Sep.", "Oct.", "Nov.", "Dic.", "Total"]
        for idx, m in enumerate(months):
            ws.cell(row=5, column=3 + idx, value=m)
        for idx, m in enumerate(months[:8]):
            ws.cell(row=5, column=16 + idx, value=m)

        # Fila 9: Total
        ws.cell(row=9, column=1, value="Total")
        ws.cell(row=9, column=3, value=2254206996.38)

        # Fila 11: Amazonas
        ws.cell(row=11, column=1, value="Amazonas")
        ws.cell(row=11, column=3, value=1391100.08)

        # Fila 27: Callao (c2)
        ws.cell(row=27, column=2, value="Provincia Constitucional del Callao")
        ws.cell(row=27, column=3, value=5000000.0)

        df = transform_sheet(ws)

        # Verificar columnas
        expected_cols = ["Departamento", "Periodo", "Anio", "Mes", "Monto_Recaudado"]
        self.assertEqual(list(df.columns), expected_cols)

        # Verificar corrección del año (2025 -> 2026 en el segundo bloque)
        years = sorted(df["Anio"].unique())
        self.assertEqual(years, [2025, 2026])

        # Verificar departamento Callao mapeado
        self.assertIn("Callao", df["Departamento"].unique())

        # Verificar valor de muestra en Total 2025_Ene.
        total_row = df[(df["Departamento"] == "Total") & (df["Periodo"] == "2025_Ene.")]
        self.assertEqual(len(total_row), 1)
        self.assertAlmostEqual(total_row.iloc[0]["Monto_Recaudado"], 2254206996.38)

    @unittest.skipIf(not TRASH_EXCEL.exists(), "Archivo real cdro_A13.xlsx no encontrado en Trash")
    def test_real_excel_transformation(self):
        """Prueba de integracion con el archivo real cdro_A13.xlsx."""
        df_monto, df_var = process_both_ingresos(str(TRASH_EXCEL))

        # Verificar cantidad de filas: 28 jurisdicciones x 260 meses = 7280
        self.assertEqual(len(df_monto), 7280)
        self.assertEqual(len(df_var), 7280)

        # Verificar muestra de Total 2005_Ene.
        monto_row = df_monto[(df_monto["Departamento"] == "Total") & (df_monto["Periodo"] == "2005_Ene.")]
        self.assertAlmostEqual(monto_row.iloc[0]["Monto_Recaudado"], 2254206996.38, places=2)

        var_row = df_var[(df_var["Departamento"] == "Total") & (df_var["Periodo"] == "2005_Ene.")]
        self.assertAlmostEqual(var_row.iloc[0]["Monto_Recaudado"], 8.84228, places=4)


if __name__ == "__main__":
    unittest.main()
